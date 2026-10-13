#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
inbox 一键解析：扫描 inbox/ 中的成员材料（Word），提取字段与照片，
归档并生成数据草稿，供助手补全后合并进 data/members.json。

用法：
    python tools/inbox_ingest.py             # 解析 + 归档 + 写草稿
    python tools/inbox_ingest.py --dry-run   # 只解析并打印，不改动任何文件

流程：
  1. 解析 docx 字段（姓名 / 教育经历 / 邮箱 / 研究领域 / 自我介绍 / 论文）
  2. 提取内嵌照片 → images/people/_pending_<姓名>.jpeg
  3. docx 归档 → member materials/BIRD Lab个人简介-<姓名>.docx
  4. 生成草稿 data/_pending_members.json（列出待补字段）

之后：助手补 slug / 英文名 / 中英翻译 / 角色 / 分组 → 合并进 data/members.json → 跑生成器。
"""
import json
import re
import sys
import zipfile
from pathlib import Path

# 部分 Word 导出的 docx，其内嵌图片的 CRC 记录有瑕疵；python-docx 读整包时会校验失败。
# 这里在解析前关闭 CRC 校验（只影响本脚本读取，不修改文件）。
zipfile.ZipExtFile._update_crc = lambda self, data: None

ROOT = Path(__file__).resolve().parent.parent
INBOX = ROOT / "inbox"
DRAFT = ROOT / "data" / "_pending_members.json"
MATERIALS = ROOT / "member materials"
PHOTOS = ROOT / "images" / "people"

LABELS = {
    "姓名": "name",
    "教育与工作经历": "education_raw",
    "教育经历": "education_raw",
    "电子邮箱": "email",
    "邮箱": "email",
    "研究领域": "research_raw",
    "研究方向": "research_raw",
    "研究领域（或感兴趣的研究方向）": "research_raw",
    "自我介绍": "bio_raw",
    "个人简介": "bio_raw",
    "论文发表情况": "publications_raw",
}
NEEDS = ["slug（拼音）", "name_en", "中英双语翻译", "role/group", "education 结构化"]


def write_lf(path, text):
    with path.open("w", encoding="utf-8", newline="\n") as f:
        f.write(text)


def norm_label(text):
    clean = re.sub(r"[（(].*?[)）]", "", text).strip().strip("：:")
    for key, field in LABELS.items():
        if clean == key or clean.startswith(key):
            return field
    return None


def extract_fields(path):
    """用 python-docx 读表格/段落，返回 {字段: 值}。"""
    try:
        import docx
    except ImportError:
        raise SystemExit("缺少 python-docx，无法解析 Word。请使用 DSH 自带的 Python 运行时。")

    doc = docx.Document(str(path))
    out = {}
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if len(cells) < 2:
                continue
            field = norm_label(cells[0])
            value = "\n".join(x for x in cells[1:] if x).strip()
            if field and value and field not in out:
                out[field] = value
    if "name" not in out:
        for p in doc.paragraphs:
            t = p.text.strip()
            if not t:
                continue
            field = norm_label(t)
            if field and field not in out:
                out[field] = ""
        # 段落兜底：第一个非空段落当作姓名
        if "name" not in out:
            texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            if texts:
                out["name"] = texts[0]
    return out


def extract_photo(path, name, dry):
    """提取内嵌的第一张图片；返回保存路径（相对仓库）。"""
    try:
        import docx
    except ImportError:
        return ""
    doc = docx.Document(str(path))
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            ext = Path(rel.target_part.partname).suffix or ".jpeg"
            target = PHOTOS / f"_pending_{name}{ext}"
            if not dry:
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_bytes(rel.target_part.blob)
            return str(target.relative_to(ROOT)).replace("\\", "/")
    return ""


def main():
    dry = "--dry-run" in sys.argv
    if not INBOX.exists():
        print("inbox/ 不存在")
        return 1
    files = [p for p in sorted(INBOX.iterdir()) if p.suffix.lower() in (".docx", ".doc") and not p.name.startswith("~$")]
    if not files:
        print("inbox/ 里没有待处理的 Word 材料。（把成员材料丢进 inbox/ 再运行本脚本）")
        return 0

    pending = []
    for i, path in enumerate(files, 1):
        fields = extract_fields(path)
        name = (fields.get("name") or path.stem).replace(" ", "")
        photo = extract_photo(path, name, dry)
        rec = {
            "source_file": path.name,
            "name_zh": name,
            "name_en": "",
            "slug": "",
            "email": fields.get("email", ""),
            "education_raw": fields.get("education_raw", ""),
            "research_raw": fields.get("research_raw", ""),
            "bio_raw": fields.get("bio_raw", ""),
            "publications_raw": fields.get("publications_raw", ""),
            "photo_pending": photo,
            "needs": NEEDS,
        }
        pending.append(rec)

        edu_lines = [l for l in rec["education_raw"].splitlines() if l.strip()]
        print(f"[{i}] {name}")
        print(f"    邮箱: {rec['email'] or '（缺）'}")
        print(f"    教育经历: {len(edu_lines)} 行")
        print(f"    研究领域: {'有' if rec['research_raw'] else '（缺）'}")
        print(f"    自我介绍: {len(rec['bio_raw'])} 字")
        print(f"    照片: {'已提取 → ' + photo if photo else '（无内嵌照片）'}")

        if not dry:
            target = MATERIALS / f"BIRD Lab个人简介-{name}.docx"
            target.parent.mkdir(parents=True, exist_ok=True)
            if target.exists():
                print(f"    [跳过归档] {target.name} 已存在")
            else:
                path.replace(target)
                print(f"    已归档 → member materials/{target.name}")

    if dry:
        print("\n（--dry-run：未改动任何文件）")
    else:
        write_lf(DRAFT, json.dumps({"pending": pending}, ensure_ascii=False, indent=2) + "\n")
        print(f"\n草稿已写入 {DRAFT.relative_to(ROOT)}（待助手补全后合并进 data/members.json）")
    print(f"待补字段：{', '.join(NEEDS)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
